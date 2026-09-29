from __future__ import annotations

from pathlib import Path
import re

import pandas as pd
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill


ROOT = Path("results") / "canales" / "canales_por_capa"
START_FOLDER = "P111"
OUTPUT_FILE = ROOT / "resumen_canales_pacientes_desde_P111.xlsx"


def natural_key(name: str) -> tuple[int, str, str]:
    match = re.match(r"^P(\d+)(?:_(derecha|izquierda))?$", name)
    if match:
        return int(match.group(1)), match.group(2) or "", name
    return 10**9, "", name


def load_patient_folders() -> list[Path]:
    folders = [path for path in ROOT.iterdir() if path.is_dir() and path.name.startswith("P")]
    folders = sorted(folders, key=lambda path: natural_key(path.name))
    started = False
    selected: list[Path] = []
    for folder in folders:
        if folder.name == START_FOLDER:
            started = True
        if started:
            selected.append(folder)
    return selected


def style_sheet(ws) -> None:
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A5"
    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 12
    ws.column_dimensions["E"].width = 16
    ws.column_dimensions["F"].width = 12
    ws.column_dimensions["G"].width = 16
    ws.column_dimensions["H"].width = 4
    for col in ["I", "J", "K", "L", "M", "N", "O", "P"]:
        ws.column_dimensions[col].width = 18


def write_dataframe(ws, df: pd.DataFrame, start_row: int = 4) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    center = Alignment(horizontal="center", vertical="center")

    for col_idx, column_name in enumerate(df.columns, start=1):
        cell = ws.cell(row=start_row, column=col_idx, value=column_name)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center

    for row_offset, row in enumerate(df.itertuples(index=False), start=1):
        for col_idx, value in enumerate(row, start=1):
            cell = ws.cell(row=start_row + row_offset, column=col_idx, value=value)
            cell.alignment = center
            if df.columns[col_idx - 1] == "dice_slice" and value is not None:
                cell.number_format = "0.000000"


def add_image(ws, image_path: Path, anchor: str, width: int = 470) -> None:
    if not image_path.exists():
        return
    img = XLImage(str(image_path))
    if img.width and img.height:
        scale = min(width / img.width, 1.0)
        img.width = int(img.width * scale)
        img.height = int(img.height * scale)
    ws.add_image(img, anchor)


def build_workbook() -> None:
    patient_folders = load_patient_folders()
    if not patient_folders:
        raise RuntimeError(f"No se encontraron carpetas de pacientes desde {START_FOLDER} en {ROOT}")

    wb = Workbook()
    wb.remove(wb.active)

    for folder in patient_folders:
        csv_path = folder / "dice_por_capa.csv"
        if not csv_path.exists():
            continue

        df = pd.read_csv(csv_path)
        ws = wb.create_sheet(title=folder.name)
        style_sheet(ws)

        ws["A1"] = f"Paciente: {folder.name}"
        ws["A1"].font = Font(bold=True, size=14)
        ws["A2"] = str(csv_path)
        ws["A2"].font = Font(italic=True, color="666666")
        ws["A3"] = "Tabla dice_por_capa y tres imágenes del paciente"
        ws["A3"].font = Font(bold=True)

        write_dataframe(ws, df, start_row=5)

        add_image(ws, folder / "inicio_conteo_canales.png", "I2")
        add_image(ws, folder / "fin_conteo_canales.png", "I24")
        add_image(ws, folder / f"{folder.name}_overlap_graph.png", "I46")

        last_row = 5 + len(df) + 2
        for row in range(5, last_row + 1):
            ws.row_dimensions[row].height = 20

    OUTPUT_FILE.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT_FILE)
    print(f"Saved: {OUTPUT_FILE}")


if __name__ == "__main__":
    build_workbook()